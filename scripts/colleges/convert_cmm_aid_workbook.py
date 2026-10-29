"""Convert the CMM "Domestic Undergraduate Need-Based and Merit Aid" workbook
into the college-aid CSV contract (docs/college-aid-csv-contract.md).

The workbook is untrusted input: run with ``python -I`` and pass its path as an
argument. Output is deterministic (sorted by state, name), so a re-run is
byte-identical.

    uv run python -I scripts/colleges/convert_cmm_aid_workbook.py aid.xlsx \
        --release 2026-09 --out OUT_DIR [--colleges-export colleges-export.csv]
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import uuid
import warnings
from pathlib import Path

# -I drops cwd and the script dir from sys.path; locate the repo from this file, never cwd.
REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

import openpyxl

from scripts.colleges.aid_workbook_values import (
    ConversionError,
    clean,
    fmt_money,
    fmt_pct,
    num,
    parse_year,
)
from src.colleges.name_key import normalize_name

HEADERS = [
    "college_id", "name", "state", "region", "control", "coa_in_state", "coa_out_of_state", "coa_year",
    "coa_estimated", "full_time_undergrads", "aid_application", "pct_need_met",
    "pct_non_need_receiving_merit", "avg_merit_award", "cds_year", "source", "release_label", "data_notes",
]
# Workbook header (whitespace-collapsed, lowercase, prefix) -> internal key. Detected by name, not position.
COLUMN_PREFIXES = {
    "institution": "name", "region": "region", "state": "state", "public/ private": "control",
    "2025-26 estimated total cost of attendance (in-state)": "coa_in_state",
    "2025-26 total estimated cost of attendance (out-of-state)": "coa_out_of_state",
    "full-time undergrads": "full_time_undergrads", "application for institutional": "aid_application",
    "average percent of need met": "pct_need_met", "percent of non-need undergrads receiving": "pct_merit",
    "average merit award to non-need": "avg_merit_award", "cds year": "cds_year", "source": "source",
    "average percent of need unmet": "j_unmet", "percent of non-need undergrads paying": "l_full",
}
COA_YEAR = "2025-26"  # workbook footnote 5: 2024-25 COA raised 4%, labelled 2025-26
LOW_NEED_MET, HIGH_MERIT_RATIO = 0.10, 0.5  # review thresholds only; values are never changed


def find_header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 15):
        if clean(ws.cell(r, 1).value).lower() == "institution":
            cols: dict[str, int] = {}
            for c in range(1, ws.max_column + 1):
                label = clean(ws.cell(r, c).value).lower()
                for prefix, key in COLUMN_PREFIXES.items():
                    if label.startswith(prefix) and key not in cols:
                        cols[key] = c
                        break
            missing = {*COLUMN_PREFIXES.values()} - cols.keys()
            if missing:
                raise ConversionError(f"workbook header is missing columns: {sorted(missing)}")
            return r, cols
    raise ConversionError("no 'Institution' header row found")


def load_ids(export: Path | None, aliases: Path) -> dict[tuple[str, str], str]:
    """(normalized name, state) -> college_id from a colleges export, then aliases (aliases win)."""
    ids: dict[tuple[str, str], str] = {}
    for path, name_col in ((export, "name"), (aliases, "source_name")):
        if not path or not path.exists():
            continue
        with path.open(newline="", encoding="utf-8") as fh:
            for rec in csv.DictReader(fh):
                if rec.get("college_id"):
                    uuid.UUID(rec["college_id"])  # validates
                    ids[(normalize_name(rec[name_col]), rec["state"].strip().upper())] = rec["college_id"]
    return ids


def convert(xlsx: Path, release: str, ids: dict) -> tuple[list[dict], list[dict]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # openpyxl: unknown extension in an Excel table
        ws = openpyxl.load_workbook(xlsx, data_only=True, read_only=False).worksheets[0]
    header_row, cols = find_header(ws)
    rows: list[dict] = []
    warns: list[dict] = []

    for r in range(header_row + 1, ws.max_row + 1):
        if not clean(ws.cell(r, cols["name"]).value):
            break  # table ends at the first blank institution (footnotes follow)
        raw = {k: ws.cell(r, c).value for k, c in cols.items()}
        name, state = clean(raw["name"]), clean(raw["state"]).upper()

        def warn(kind: str, detail: str, name=name, state=state, r=r) -> None:
            warns.append({"name": name, "state": state, "source_row": r, "kind": kind, "detail": detail})

        if clean(raw["name"]) != str(raw["name"]):
            warn("data_fix", "whitespace trimmed in name")
        year, note = parse_year(raw["cds_year"], r)
        if note:
            warn("cds_year_not_reported", note)
        elif clean(raw["cds_year"]) != str(raw["cds_year"]):
            warn("data_fix", "whitespace trimmed in CDS year")

        need = num(raw["pct_need_met"], "pct_need_met", r)
        merit = num(raw["pct_merit"], "pct_non_need_receiving_merit", r)
        if any(v is not None and not 0 <= v <= 1 for v in (need, merit)):
            raise ConversionError(f"row {r}: percent outside 0..1: {need}, {merit}")
        award = num(raw["avg_merit_award"], "avg_merit_award", r)
        coa_in, coa_out = (num(raw[k], k, r) for k in ("coa_in_state", "coa_out_of_state"))
        fte = num(raw["full_time_undergrads"], "full_time_undergrads", r)
        control = clean(raw["control"]).lower()
        if control not in ("public", "private"):
            raise ConversionError(f"row {r}: invalid control {raw['control']!r}")
        # Values below are kept as given; these flags only ask CMM to review.

        if control == "public" and coa_out is None:
            warn("missing_value", "public with no out-of-state COA")
        if need is not None and need < LOW_NEED_MET:
            warn("outlier", f"pct_need_met {fmt_pct(need)} is very low; kept as given")
        if merit == 1 and award == 0:
            warn("contradiction", "100% receive merit but average award is 0")
        if merit == 0 and award is None:
            warn("contradiction", "0% receive merit and average award is N/R")
        if award and (coa_out or coa_in) and award / (coa_out or coa_in) > HIGH_MERIT_RATIO:
            warn("outlier", "average merit award is over half of COA; read with the merit percent")
        # Formula columns are dropped; flag any that disagree with the inputs they were derived from.
        for derived, base, key in ((raw["j_unmet"], need, "pct_need_unmet"), (raw["l_full"], merit, "pct_full_coa")):
            if isinstance(derived, (int, float)) and base is not None and abs(derived - (1 - base)) > 0.0005:
                warn("derived_mismatch", f"{key} in workbook disagrees with 1 - input")

        has_coa = coa_in is not None or coa_out is not None
        rows.append({
            "college_id": ids.get((normalize_name(name), state), ""), "name": name, "state": state,
            "region": clean(raw["region"]), "control": control,
            "coa_in_state": fmt_money(coa_in), "coa_out_of_state": fmt_money(coa_out),
            "coa_year": COA_YEAR if has_coa else "", "coa_estimated": "true" if has_coa else "false",
            "full_time_undergrads": "" if fte is None else str(int(fte)),
            "aid_application": clean(raw["aid_application"]),
            "pct_need_met": fmt_pct(need), "pct_non_need_receiving_merit": fmt_pct(merit),
            "avg_merit_award": fmt_money(award), "cds_year": year, "source": clean(raw["source"]),
            "release_label": release, "data_notes": note,
        })

    keys = [(normalize_name(x["name"]), x["state"]) for x in rows]
    if len(set(keys)) != len(keys):
        dupes = sorted({k for k in keys if keys.count(k) > 1})
        raise ConversionError(f"duplicate (name, state) in workbook: {dupes}")
    with_id = [(x["college_id"], x["cds_year"]) for x in rows if x["college_id"]]
    if len(set(with_id)) != len(with_id):
        raise ConversionError("duplicate (college_id, cds_year): two rows resolve to one college")
    rows.sort(key=lambda x: (x["state"], x["name"]))
    warns.sort(key=lambda w: (w["state"], w["name"], w["kind"], w["detail"]))
    return rows, warns


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description="Convert the CMM aid workbook to the college-aid CSV contract")
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--release", required=True, help="release label YYYY-MM")
    ap.add_argument("--out", required=True, type=Path, help="output directory (outside the repo)")
    ap.add_argument("--colleges-export", type=Path, help="colleges-export.csv from the admin list")
    ap.add_argument("--aliases", type=Path, default=Path(__file__).with_name("college_aliases.csv"))
    args = ap.parse_args()
    try:
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.release):
            raise ConversionError("--release must be YYYY-MM")
        rows, warns = convert(args.workbook, args.release, load_ids(args.colleges_export, args.aliases))
    except (ConversionError, ValueError, OSError, KeyError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    args.out.mkdir(parents=True, exist_ok=True)
    write_csv(args.out / f"college-aid-{args.release}.csv", HEADERS, rows)
    write_csv(args.out / "review-warnings.csv", ["name", "state", "source_row", "kind", "detail"], warns)
    print(f"{len(rows)} rows, {len(warns)} warnings -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
